#!/usr/bin/env python3
"""Ce qu'on trouve derrière le QR code, montré à l'écran plutôt que décrit.

La diapo 12 de la présentation dit ce que le programme contient. Celle-ci le
montre : deux captures d'écran du téléphone — la salle qui s'allume sur le
plan, et le bouton « Repérer ». Elle se place juste après la diapo du QR code.

Deux mises en page sont produites, parce que la contrainte n'est pas la place
sur la diapo mais la lisibilité depuis le fond d'un auditorium de 254 places :

  diapo-en-main-1-diapo.pptx    les deux captures côte à côte
  diapo-en-main-2-diapos.pptx   une capture par diapo, à pleine hauteur

La hauteur est le facteur limitant dans les deux cas : sous le titre il reste
10,5 cm. À deux captures, chacune perd la place du cartouche ; à une capture
par diapo, le cartouche passe à gauche et la capture prend toute la hauteur.

    python3 outils/affiches/generer-diapo-en-main.py

Fichiers séparés, et non insertion dans le générateur de la présentation : la
présentation est reprise à la main (les noms, les lots). Régénérer les quinze
diapos écraserait ces reprises, et insérer une diapo décalerait la numérotation
du pied de page. Un fichier de une ou deux diapos s'importe dans une
présentation déjà travaillée sans rien perdre. D'où, aussi, l'absence de numéro
de page : on ne connaît pas encore la pagination d'accueil.

Les captures sont figées dans ressources/ : elles datent d'un état de la page
où les salles étaient déjà publiées. Si les affectations changent, il faut
refaire les captures — relancer ce script ne suffit pas.
"""

import importlib.util
import pathlib
import subprocess
import sys

from pptx import Presentation
from pptx.util import Cm
from pptx.enum.text import PP_ALIGN

ICI = pathlib.Path(__file__).resolve().parent
RES = ICI / 'ressources'
CAPTURE_PLAN = RES / 'capture-plan.png'
CAPTURE_REPERER = RES / 'capture-reperer.png'

# Rapports largeur/hauteur relevés à l'enregistrement des captures. Écrits ici
# pour que la composition ne dépende pas de Pillow au moment du tracé.
RAPPORT_PLAN = 1080 / 1242
RAPPORT_REPERER = 1170 / 1016

HAUT_UTILE = Cm(7.15)          # sous le titre
BAS_UTILE = Cm(17.65)          # au-dessus du pied de page


def presentation():
    """Le module voisin porte un tiret dans son nom : il ne s'importe pas.
    On le charge par importlib plutôt que de recopier le bandeau, la
    signature, le pied et la palette — qui doivent rester identiques."""
    chemin = ICI / 'generer-presentation.py'
    spec = importlib.util.spec_from_file_location('generer_presentation', chemin)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def cartouche(P, diapo, x, y, l, numero, titre, lignes):
    """Le numéro, l'intitulé, puis ce qu'il faut en dire."""
    P.bloc(diapo, x, y + Cm(0.05), Cm(0.95), Cm(0.95), P.TEAL_FOND)
    P.texte(diapo, x, y + Cm(0.16), Cm(0.95), Cm(0.7),
            [P.ligne(numero, P.CONDENSEE, 18, P.BLANC, True, 1.0)],
            align=PP_ALIGN.CENTER)
    P.texte(diapo, x + Cm(1.3), y, l - Cm(1.3), Cm(1.1),
            [P.ligne(titre.upper(), P.CONDENSEE, 21, P.ENCRE, True, 1.0)])
    if lignes:
        P.texte(diapo, x + Cm(1.3), y + Cm(1.3), l - Cm(1.3),
                BAS_UTILE - y - Cm(1.3),
                [P.ligne(contenu, P.COURANTE, 12.5,
                         P.ENCRE if gras else P.GRIS, gras, 1.4)
                 for contenu, gras in lignes])


def capture(diapo, image, x, y, haut, rapport):
    large = int(haut * rapport)
    diapo.shapes.add_picture(str(image), x, y, large, haut)
    return large


# ── Une diapo : les deux captures côte à côte ──────────────────────────

def composer_une(P, diapo, prog):
    P.chrome(diapo, prog, None)
    P.titre_diapo(diapo, 'Le programme en main')

    haut_cartouche = Cm(2.5)
    # De l'air sous les captures : collées au pied, elles donnaient
    # l'impression d'être coupées.
    haut_image = BAS_UTILE - HAUT_UTILE - haut_cartouche - Cm(0.35)

    volets = [
        ('1', 'Touchez une salle',
         [("Le plan s'ouvre et votre salle s'allume.", False)],
         CAPTURE_PLAN, RAPPORT_PLAN),
        ('2', 'Touchez Repérer',
         [('Cinq kiosques au plus, sur votre téléphone.', False)],
         CAPTURE_REPERER, RAPPORT_REPERER),
    ]
    larges = [int(haut_image * r) for *_, r in volets]
    # Chaque colonne fait la largeur de sa capture, sans descendre sous la
    # place qu'il faut à l'intitulé. L'ensemble est centré : les deux captures
    # n'ont pas le même format, c'est l'axe commun qui les tient ensemble.
    colonnes = [max(l, int(Cm(10.6))) for l in larges]
    ecart = Cm(1.8)
    x = (P.LARGEUR - sum(colonnes) - ecart) // 2

    for (numero, titre, lignes, image, rapport), colonne, large in zip(
            volets, colonnes, larges):
        cartouche(P, diapo, x, HAUT_UTILE, colonne, numero, titre, lignes)
        capture(diapo, image, x + (colonne - large) // 2,
                HAUT_UTILE + haut_cartouche, haut_image, rapport)
        x += colonne + ecart


# ── Deux diapos : une capture par diapo, à pleine hauteur ──────────────

def composer_pleine(P, diapo, prog, titre_page, image, rapport,
                    numero, titre, lignes):
    P.chrome(diapo, prog, None)
    P.titre_diapo(diapo, titre_page)
    haut = BAS_UTILE - HAUT_UTILE
    large = int(haut * rapport)
    capture(diapo, image, P.LARGEUR - P.MARGE - large, HAUT_UTILE, haut, rapport)
    # L'intitulé et ses lignes forment un bloc ; on le centre en hauteur face
    # à la capture au lieu de le poser en haut d'une colonne presque vide.
    bloc_haut = Cm(1.3) + len(lignes) * Cm(0.62)
    cartouche(P, diapo, P.MARGE, HAUT_UTILE + (haut - bloc_haut) // 2,
              P.LARGEUR - 2 * P.MARGE - large - Cm(1.5),
              numero, titre, lignes)


def composer_plan(P, diapo, prog):
    composer_pleine(
        P, diapo, prog, 'Trouver sa salle', CAPTURE_PLAN, RAPPORT_PLAN,
        '1', 'Touchez une salle', [
            ('Sur chaque kiosque, le nom de la salle est un bouton.', False),
            ("Le plan du rez-de-jardin s'ouvre, votre salle s'allume et une "
             'onde la signale.', False),
            ('Les salles sont publiées le jour même : les affectations '
             'tombent tard.', False),
        ])


def composer_reperer(P, diapo, prog):
    composer_pleine(
        P, diapo, prog, 'Marquer ses kiosques', CAPTURE_REPERER, RAPPORT_REPERER,
        '2', 'Touchez Repérer', [
            ('Le bouton passe à « Repéré » et la carte prend un liseré.', False),
            ('Un compteur suit votre sélection : cinq kiosques au plus, '
             'un par rotation.', False),
            ('Conservé sur votre téléphone, y compris après la mise à jour '
             'des salles.', False),
            ("Cela ne réserve pas de place : l'accès aux kiosques se fait "
             'sur place.', True),
        ])


# ── Assemblage et contrôle ─────────────────────────────────────────────

VERSIONS = [
    ('diapo-en-main-1-diapo.pptx', [composer_une]),
    ('diapo-en-main-2-diapos.pptx', [composer_plan, composer_reperer]),
]


def construire(nom, composeurs):
    P = presentation()
    prog = P.lire_programme()
    prez = Presentation()
    prez.slide_width, prez.slide_height = P.LARGEUR, P.HAUTEUR
    for composer in composeurs:
        composer(P, prez.slides.add_slide(prez.slide_layouts[6]), prog)
    chemin = ICI / nom
    prez.save(chemin)
    return chemin


def en_pdf(pptx):
    subprocess.run(
        ['soffice', '--headless', '-env:UserInstallation=file:///tmp/lo-affiches',
         '--convert-to', 'pdf', '--outdir', str(pptx.parent), str(pptx)],
        check=True, capture_output=True, timeout=600)
    return pptx.with_suffix('.pdf')


def verifier(pdf, attendues):
    """Relit le PDF produit : le nombre de pages, le format, et tout bloc de
    texte qui descendrait sur le pied de page — c'est ce qui a cassé deux fois
    sur la présentation."""
    import pymupdf
    doc = pymupdf.open(pdf)
    anomalies = []
    if doc.page_count != attendues:
        anomalies.append(f'{doc.page_count} pages au lieu de {attendues}.')
    for page in doc:
        l, h = page.rect.width, page.rect.height
        if abs(l / h - 16 / 9) > 0.01:
            anomalies.append(f"format {l:.0f}×{h:.0f} — ce n'est pas du 16:9.")
        limite = h * (19.05 - 1.35) / 19.05
        cadres = [rect for xref in {i[0] for i in page.get_images(full=True)}
                  for rect in page.get_image_rects(xref)]
        for b in page.get_text('blocks'):
            bloc_ = pymupdf.Rect(b[:4])
            libelle = b[4].strip().replace('\n', ' ')[:40]
            if b[1] < limite and b[3] > limite + 2:
                anomalies.append(f'« {libelle} » déborde sur le pied.')
            # Une ligne de trop dans un cartouche passe derrière la capture
            # d'à côté : invisible sur la diapo, invisible au PDF si on ne
            # compare pas les encombrements. Deux points carrés de
            # recouvrement suffisent à le dire.
            for cadre in cadres:
                if (bloc_ & cadre).get_area() > 2:
                    anomalies.append(f'« {libelle} » passe sous une capture.')
                    break
    doc.close()
    return anomalies


def main():
    for image in (CAPTURE_PLAN, CAPTURE_REPERER):
        if not image.exists():
            sys.exit(f'capture manquante : {image}')
    fautes = 0
    for nom, composeurs in VERSIONS:
        pptx = construire(nom, composeurs)
        anomalies = verifier(en_pdf(pptx), len(composeurs))
        for a in anomalies:
            print(f'  ✗ {pptx.name} : {a}')
        fautes += len(anomalies)
        if not anomalies:
            print(f'✓ {pptx.name} — {len(composeurs)} diapo(s) 16:9.')
    if fautes:
        sys.exit(1)


if __name__ == '__main__':
    main()
