#!/usr/bin/env python3
"""Le programme entier sur un seul écran, pour l'espace hospitalité.

Ce n'est pas la diapo 3 en plus dense. C'est un autre problème : la diapo 3 est
projetée dans un auditorium et se lit depuis le fond, donc peu de mots et gros ;
cet écran est regardé à deux ou trois mètres par quelqu'un qui passe, et il doit
se suffire à lui-même — personne ne s'arrêtera pour scanner un QR code.

D'où la densité assumée : les trois conférences avec leur accroche, les dix
kiosques avec leur phrase de présentation **et leur salle**. Ce sont les deux
informations qui manquaient à l'écran précédent (`generer-ecrans.py`), et sans
lesquelles quelqu'un qui découvre le programme ici doit aller chercher ailleurs.

    python3 outils/affiches/generer-ecran-complet.py

Le contrôle relit le PDF et refuse tout texte qui sort de sa cellule. C'est le
seul risque réel d'une composition à cette densité : un pitch qui passe de deux
à trois lignes ne casse rien de visible dans le code, et déborde sur le voisin.
"""

import importlib.util
import pathlib
import subprocess
import sys

from pptx import Presentation
from pptx.util import Cm
from pptx.enum.text import PP_ALIGN

ICI = pathlib.Path(__file__).resolve().parent
RACINE = ICI.parents[1]
QR = RACINE / 'qr' / 'testing-event-2026.png'

LARGEUR, HAUTEUR = Cm(33.87), Cm(19.05)
BAS_BANDEAU, BAS_SIGNATURE = Cm(2.35), Cm(3.20)
HAUT_CONTENU, BAS_CONTENU = Cm(3.45), Cm(18.10)
COL_G, LARGE_G = Cm(1.25), Cm(10.30)
COL_D, LARGE_D = Cm(12.15), Cm(20.45)
HAUT_QR = Cm(1.95)


def presentation():
    """Le module voisin porte un tiret : il ne s'importe pas. On le charge par
    importlib pour reprendre la palette, les helpers de tracé et le lecteur du
    programme — une seule source pour les couleurs et pour les données."""
    spec = importlib.util.spec_from_file_location(
        'generer_presentation', ICI / 'generer-presentation.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class Feuille:
    """La diapo en cours, et la liste des cellules où du texte a le droit de
    se trouver. Le contrôle s'en sert pour dire si quelque chose déborde."""

    def __init__(self, P, diapo):
        self.P, self.diapo, self.cellules = P, diapo, []

    def cellule(self, x, y, l, h, fond, liseré=None, large_liseré=Cm(0.2)):
        self.P.bloc(self.diapo, x, y, l, h, fond)
        if liseré:
            self.P.bloc(self.diapo, x, y, large_liseré, h, liseré)
        self.cellules.append((x, y, l, h))
        return x, y, l, h


def _pastille(F, x, y, l, h, libelle, fond, encre, taille=8):
    P = F.P
    F.cellule(x, y, l, h, fond)
    P.texte(F.diapo, x, y + (h - Cm(0.42)) // 2, l, Cm(0.5),
            [P.ligne(libelle, P.CONDENSEE, taille, encre, True, 1.0)],
            align=PP_ALIGN.CENTER)


def _bandeau_horaire(F, x, y, l, h, heure, libelle, fond_heure, fond, encre):
    """Un créneau simple : pavé d'heure plein puis libellé sur fond teinté,
    comme la page dessine ses « special-slot »."""
    P = F.P
    large_heure = Cm(2.3)
    P.bloc(F.diapo, x, y, large_heure, h, fond_heure)
    F.cellule(x + large_heure, y, l - large_heure, h, fond)
    F.cellules[-1] = (x, y, l, h)
    P.texte(F.diapo, x, y + (h - Cm(0.46)) // 2, large_heure, Cm(0.55),
            [P.ligne(heure, P.CONDENSEE, 13, P.BLANC, True, 1.0)],
            align=PP_ALIGN.CENTER)
    P.texte(F.diapo, x + large_heure + Cm(0.45), y + (h - Cm(0.44)) // 2,
            l - large_heure - Cm(0.8), Cm(0.55),
            [P.ligne(libelle, P.COURANTE, 11, encre, True, 1.0)])


# Largeur moyenne d'un caractère, en centimètres par point de corps. Relevée
# sur un rendu : à 14 pt dans 9,35 cm, un titre de 113 signes tient sur trois
# lignes. Volontairement un peu pessimiste — mieux vaut une carte trop haute
# qu'un titre qui vient écrire sur son accroche.
LARGEUR_SIGNE = 0.0145

# Un point vaut 2,54/72 cm. Et l'interligne donné en nombre à python-pptx
# multiplie la hauteur naturelle de la police — environ 1,2 fois le corps —
# et non le corps lui-même. Sans ce facteur, chaque ligne était sous-estimée
# d'un cinquième et le titre de la conférence 2 venait écrire sur sa propre
# accroche. Deux corrections successives sur le même calcul : d'abord une
# valeur devinée, puis une formule à laquelle il manquait ce facteur.
POINT_EN_CM = 2.54 / 72
INTERLIGNE_NATUREL = 1.2

TAILLE_TITRE, TAILLE_CITATION = 13, 8.5


def _haut_ligne(taille, interligne):
    return Cm(taille * INTERLIGNE_NATUREL * interligne * POINT_EN_CM)


def _lignes(texte_, largeur_emu, taille):
    signes = max(1, int((largeur_emu / 360000) / (LARGEUR_SIGNE * taille)))
    return max(1, -(-len(texte_) // signes))


def _haut_seance(titre, citation, l):
    """Ce qu'il faut à une séance : l'entête, le titre, l'accroche, les marges."""
    tl = l - Cm(0.95)
    return (Cm(0.24) + Cm(0.40)
            + _lignes(titre, tl, TAILLE_TITRE) * _haut_ligne(TAILLE_TITRE, 1.06) + Cm(0.10)
            + _lignes(citation, tl, TAILLE_CITATION) * _haut_ligne(TAILLE_CITATION, 1.25) + Cm(0.28))


def _seance(F, x, y, l, h, heure, genre, titre, citation, accent, encre_genre):
    P = F.P
    F.cellule(x, y, l, h, P.BLANC, accent)
    tx, tl = x + Cm(0.55), l - Cm(0.95)
    haut_cit = _lignes(citation, tl, TAILLE_CITATION) * _haut_ligne(TAILLE_CITATION, 1.25) + Cm(0.08)
    P.texte(F.diapo, tx, y + Cm(0.24), tl, Cm(0.42),
            [P.ligne(f'{heure} · {genre.upper()}', P.CONDENSEE, 10,
                     encre_genre, True, 1.0)])
    P.texte(F.diapo, tx, y + Cm(0.68), tl, h - Cm(0.68) - haut_cit - Cm(0.25),
            [P.ligne(titre, P.CONDENSEE, TAILLE_TITRE, P.ENCRE, True, 1.06)])
    P.texte(F.diapo, tx, y + h - haut_cit - Cm(0.22), tl, haut_cit,
            [P.ligne(citation, P.COURANTE, TAILLE_CITATION, P.GRIS, interligne=1.25)])


def _kiosque(F, x, y, l, h, k):
    P = F.P
    fond = P.TEAL_FOND if int(k['numero']) % 2 else P.ROUGE_FONCE
    large_num = Cm(1.1)
    P.bloc(F.diapo, x, y, large_num, h, fond)
    F.cellule(x + large_num, y, l - large_num, h, P.BLANC)
    F.cellules[-1] = (x, y, l, h)
    P.texte(F.diapo, x, y + (h - Cm(0.6)) // 2, large_num, Cm(0.7),
            [P.ligne(k['numero'], P.CONDENSEE, 16, P.BLANC, True, 1.0)],
            align=PP_ALIGN.CENTER)

    tx = x + large_num + Cm(0.45)
    tl = l - large_num - Cm(0.85)
    large_salle = Cm(2.75)
    P.texte(F.diapo, tx, y + Cm(0.24), tl - large_salle - Cm(0.2), Cm(0.45),
            [P.ligne(k['titre'], P.CONDENSEE, 13, P.ENCRE, True, 1.0)])
    _pastille(F, x + l - Cm(0.4) - large_salle, y + Cm(0.2), large_salle, Cm(0.52),
              k['salle'], P.CYAN_CLAIR, P.TEAL, 9)
    P.texte(F.diapo, tx, y + Cm(0.82), tl, Cm(0.7),
            [P.ligne(k['pitch'], P.COURANTE, 8, P.GRIS, interligne=1.22)])


def composer(P, diapo, prog):
    F = Feuille(P, diapo)
    m, s, r = prog['moments'], prog['seances'], prog['rotations']

    # ── Chrome ────────────────────────────────────────────────────────
    P.bloc(diapo, 0, 0, LARGEUR, HAUTEUR, P.GRIS_CLAIR)
    P.bloc(diapo, 0, 0, LARGEUR, BAS_BANDEAU, P.ENCRE)
    diapo.shapes.add_picture(str(P.LOGO), COL_G, Cm(0.5),
                             Cm(6.2), Cm(6.2 * 104 / 480))
    F.cellules.append((0, 0, LARGEUR, BAS_SIGNATURE))
    P.texte(diapo, Cm(9.0), Cm(0.42), Cm(14.0), Cm(1.6), [
        P.ligne('PROGRAMME DE LA JOURNÉE', P.CONDENSEE, 22, P.BLANC, True, 1.0),
        P.ligne('Lundi 14 septembre 2026 · Auditorium Seine & kiosques '
                'du rez-de-jardin', P.COURANTE, 10, P.BLEU_CLAIR, interligne=1.4),
    ])
    P.texte(diapo, Cm(23.5), Cm(0.5), Cm(9.1), Cm(1.5), [
        P.ligne('BUSINESS CENTER CAA', P.CONDENSEE, 15, P.CYAN, True, 1.0),
        P.ligne('36/44 bd de Vaugirard · Paris', P.COURANTE, 9.5, P.BLEU_CLAIR,
                interligne=1.4),
    ], align=PP_ALIGN.RIGHT)

    P.bloc(diapo, 0, BAS_BANDEAU, LARGEUR, BAS_SIGNATURE - BAS_BANDEAU, P.ROUGE)
    P.texte(diapo, COL_G, BAS_BANDEAU + Cm(0.16), LARGEUR - 2 * COL_G, Cm(0.6),
            [P.ligne('★  ' + prog['signature'].upper() + '  ★',
                     P.CONDENSEE, 13, P.BLANC, True, 1.0)], align=PP_ALIGN.CENTER)

    P.bloc(diapo, 0, BAS_CONTENU, LARGEUR, HAUTEUR - BAS_CONTENU, P.ENCRE)
    F.cellules.append((0, BAS_CONTENU, LARGEUR, HAUTEUR - BAS_CONTENU))
    P.texte(diapo, COL_G, BAS_CONTENU + Cm(0.26), Cm(20), Cm(0.5),
            [P.ligne('Testing Event · Crédit Agricole Assurances · '
                     '14 septembre 2026', P.COURANTE, 9, P.BLEU_CLAIR, interligne=1.0)])
    P.texte(diapo, Cm(20), BAS_CONTENU + Cm(0.26), Cm(12.6), Cm(0.5),
            [P.ligne("20 min d'atelier · 5 min de questions · "
                     '5 min pour changer de salle', P.COURANTE, 9, P.CYAN,
                     interligne=1.0)], align=PP_ALIGN.RIGHT)

    # ── Intitulés de colonne ──────────────────────────────────────────
    for x, l, intitule, couleur in (
            (COL_G, LARGE_G, 'MATIN — CONFÉRENCES', P.TEAL),
            (COL_D, LARGE_D, 'APRÈS-MIDI — KIOSQUES & ATELIERS', P.ROUGE_FONCE)):
        P.texte(diapo, x, HAUT_CONTENU, l, Cm(0.75),
                [P.ligne(intitule, P.CONDENSEE, 15, couleur, True, 1.0)])
        F.cellules.append((x, HAUT_CONTENU, l, Cm(0.75)))

    y0 = HAUT_CONTENU + Cm(0.90)

    # ── Colonne du matin ──────────────────────────────────────────────
    haut_simple, ecart = Cm(0.80), Cm(0.14)
    # Les trois séances se partagent ce qui reste, mais pas à parts égales :
    # le titre de la conférence 2 fait le double des autres. On mesure ce
    # qu'il faut à chacune, puis on répartit le reliquat.
    dispo = (BAS_CONTENU - HAUT_QR - Cm(0.26)) - y0 - 3 * haut_simple - 5 * ecart
    besoins = [_haut_seance(x['titre'], x['citation'], LARGE_G) for x in s]
    rab = (dispo - sum(besoins)) // len(besoins)
    hauteurs = [b + rab for b in besoins]
    y = y0
    _bandeau_horaire(F, COL_G, y, LARGE_G, haut_simple, P.hhmm(540), m[540][1],
                     P.TEAL_FOND, P.CYAN_CLAIR, P.TEAL)
    y += haut_simple + ecart
    _bandeau_horaire(F, COL_G, y, LARGE_G, haut_simple, P.hhmm(585), m[585][1],
                     P.ENCRE, P.LILAS, P.ENCRE)
    y += haut_simple + ecart
    accents = [(P.CYAN, P.TEAL), (P.ROUGE, P.ROUGE_FONCE),
               (P.VERT_PRAIRIE, P.VERT_FONCE)]
    for seance, (accent, encre_genre), haut in zip(s, accents, hauteurs):
        _seance(F, COL_G, y, LARGE_G, haut, P.hhmm(seance['debut']),
                seance['genre'], seance['titre'], seance['citation'],
                accent, encre_genre)
        y += haut + ecart
    _bandeau_horaire(F, COL_G, y, LARGE_G, haut_simple, P.hhmm(735), m[735][1],
                     P.ROUGE_FONCE, P.ORANGE_CLAIR, P.BRUN)

    # Le QR code ferme la colonne. L'écran se suffit à lui-même — le code n'est
    # là que pour qui voudra la version qui suit les changements de salle.
    cote = HAUT_QR
    y = BAS_CONTENU - cote
    F.cellule(COL_G, y, LARGE_G, cote, P.BLANC)
    P.bloc(diapo, COL_G + Cm(0.12), y + Cm(0.12), cote - Cm(0.24),
           cote - Cm(0.24), P.BLANC)
    diapo.shapes.add_picture(str(QR), COL_G + Cm(0.18), y + Cm(0.18),
                             cote - Cm(0.36), cote - Cm(0.36))
    P.texte(diapo, COL_G + cote + Cm(0.35), y + Cm(0.42),
            LARGE_G - cote - Cm(0.7), Cm(1.2), [
                P.ligne('LE PROGRAMME SUR VOTRE TÉLÉPHONE',
                        P.CONDENSEE, 11, P.ENCRE, True, 1.0),
                P.ligne('Le plan des salles, à jour.',
                        P.COURANTE, 8.5, P.GRIS, interligne=1.35),
            ])

    # ── Colonne de l'après-midi ───────────────────────────────────────
    y = y0
    _bandeau_horaire(F, COL_D, y, LARGE_D, haut_simple, P.hhmm(820), m[820][1],
                     P.ENCRE, P.LILAS, P.ENCRE)
    y += haut_simple + Cm(0.20)

    F.cellule(COL_D, y, LARGE_D, haut_simple, P.ENCRE)
    P.texte(diapo, COL_D + Cm(0.55), y + Cm(0.24), LARGE_D - Cm(1.1), Cm(0.5),
            [P.ligne(f"{P.hhmm(r[0][0])} – {P.hhmm(r[-1][1])}  ·  "
                     '5 ROTATIONS DE 30 MIN  ·  10 KIOSQUES EN SIMULTANÉ  ·  '
                     'CHOISISSEZ VOS SESSIONS', P.CONDENSEE, 13, P.CYAN, True, 1.0)])
    y += haut_simple + Cm(0.20)

    haut_k, ecart_k = Cm(1.97), Cm(0.13)
    large_k = (LARGE_D - Cm(0.31)) // 2
    for i, k in enumerate(prog['kiosques']):
        _kiosque(F, COL_D + int((i % 2) * (large_k + Cm(0.31))),
                 y + int((i // 2) * (haut_k + ecart_k)), large_k, haut_k, k)
    y += 5 * haut_k + 4 * ecart_k + Cm(0.20)

    _bandeau_horaire(F, COL_D, y, LARGE_D, haut_simple, P.hhmm(990), m[990][1],
                     P.TEAL_FOND, P.CYAN_CLAIR, P.TEAL)
    return F.cellules


def construire(prog, cellules_vues):
    P = presentation()
    prez = Presentation()
    prez.slide_width, prez.slide_height = LARGEUR, HAUTEUR
    diapo = prez.slides.add_slide(prez.slide_layouts[6])
    cellules_vues.extend(composer(P, diapo, prog))
    chemin = ICI / 'ecran-programme-complet.pptx'
    prez.save(chemin)
    return chemin


def en_pdf(pptx):
    subprocess.run(
        ['soffice', '--headless', '-env:UserInstallation=file:///tmp/lo-affiches',
         '--convert-to', 'pdf', '--outdir', str(pptx.parent), str(pptx)],
        check=True, capture_output=True, timeout=600)
    return pptx.with_suffix('.pdf')


def verifier(pdf, cellules, marge=0.12):
    """Chaque bloc de texte doit tenir dans une cellule déclarée.

    À cette densité, le seul défaut qui compte est le débordement : un pitch
    qui passe de deux à trois lignes, un titre qui prend une ligne de plus.
    Rien ne le signale dans le code, et à l'écran il recouvre le voisin. On
    compare donc les encombrements du PDF aux cellules dessinées.
    """
    import pymupdf
    cm = 72 / 2.54
    boites = [(x / 360000, y / 360000, (x + l) / 360000, (y + h) / 360000)
              for x, y, l, h in cellules]
    doc = pymupdf.open(pdf)
    anomalies = []
    if doc.page_count != 1:
        anomalies.append(f'{doc.page_count} pages au lieu d\'une.')
    page = doc[0]
    if (abs(page.rect.width / cm - 33.87) > 0.05
            or abs(page.rect.height / cm - 19.05) > 0.05):
        anomalies.append(f'{page.rect.width / cm:.2f} × {page.rect.height / cm:.2f} cm, '
                         '16:9 attendu.')
    # Au fragment, pas au bloc : l'extracteur regroupe volontiers deux textes
    # distants posés sur la même ligne de base — les deux intitulés de colonne
    # arrivaient fusionnés en un pavé large de 18 cm, qui ne correspondait à
    # rien de dessiné.
    frags = []
    for bloc_ in page.get_text('dict')['blocks']:
        for ligne_ in bloc_.get('lines', []):
            for frag in ligne_['spans']:
                contenu = frag['text'].strip()
                if not contenu:
                    continue
                x0, y0, x1, y1 = (v / cm for v in frag['bbox'])
                frags.append(((x0, y0, x1, y1), contenu))
                if any(x0 >= bx0 - marge and y0 >= by0 - marge
                       and x1 <= bx1 + marge and y1 <= by1 + marge
                       for bx0, by0, bx1, by1 in boites):
                    continue
                anomalies.append(f'« {contenu[:52]} » sort de sa cellule '
                                 f'({x0:.1f},{y0:.1f})→({x1:.1f},{y1:.1f}).')
    # Deuxième invariant, indépendant du premier : deux fragments ne doivent
    # jamais se recouvrir. Le titre trop long qui vient écrire par-dessus sa
    # propre accroche reste, lui, à l'intérieur de sa cellule — le contrôle
    # de débordement ne peut pas le voir.
    for i, (a, ta) in enumerate(frags):
        for b, tb in frags[i + 1:]:
            largeur = min(a[2], b[2]) - max(a[0], b[0])
            hauteur = min(a[3], b[3]) - max(a[1], b[1])
            if largeur <= 0 or hauteur <= 0:
                continue
            aire = largeur * hauteur
            petit = min((a[2] - a[0]) * (a[3] - a[1]),
                        (b[2] - b[0]) * (b[3] - b[1]))
            if petit > 0 and aire / petit > 0.18:
                anomalies.append(f'« {ta[:34]} » et « {tb[:34]} » se recouvrent.')
    doc.close()
    return anomalies


def main():
    P = presentation()
    prog = P.lire_programme()
    manquantes = [k['numero'] for k in prog['kiosques']
                  if k['salle'] == 'Salle à confirmer']
    if manquantes:
        print('  ! salles non publiées pour les kiosques '
              + ', '.join(manquantes))
    cellules = []
    pptx = construire(prog, cellules)
    anomalies = verifier(en_pdf(pptx), cellules)
    for a in anomalies:
        print('  ✗', a)
    if anomalies:
        sys.exit(1)
    print(f'✓ {pptx.name} — un écran 16:9, {len(cellules)} cellules, '
          'rien qui déborde.')


if __name__ == '__main__':
    main()
